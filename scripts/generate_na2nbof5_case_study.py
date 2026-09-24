from __future__ import annotations

import csv
import json
from pathlib import Path
from textwrap import dedent

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "case_study_na2nbof5"
FIG = OUT / "visual_assets"
OUT.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)
DATE = "2026-09-23"

NAVY, TEAL, PALE, INK, AMBER, RED, WHITE = "16324F", "2A7F83", "E8F2F2", "1F2933", "C58A19", "A64040", "FFFFFF"

SOURCES = {
    "primary": ("Morkhova et al. (2026), A new sodium-ion conductor Na2NbOF5", "https://doi.org/10.1016/j.jpcs.2026.113889"),
    "primary_preprint": ("Morkhova et al., SSRN preprint record", "https://doi.org/10.2139/ssrn.6481875"),
    "screening": ("Hierarchical screening of ICSD ... mixed polyanionic oxohalides", "https://doi.org/10.1007/s10008-025-06205-4"),
    "phosphor": ("Na2NbOF5:Mn4+ phosphor record", "https://pubmed.ncbi.nlm.nih.gov/34576541/"),
    "he": ("He, Chen & Lai, Na diffusion with MLIPs", "https://doi.org/10.2139/ssrn.4431516"),
    "miyagawa": ("Miyagawa et al., MLMD ion migration", "https://arxiv.org/abs/2401.11244"),
    "benchmark": ("Aghoghovbia et al., Probing MLIPs on ion transport", "https://doi.org/10.1002/aidi.70143"),
    "horizons": ("ML-assisted discovery of sodium superionic conductors", "https://doi.org/10.1039/D5MH01176K"),
    "deepmd": ("Neural-network potential for solid-state electrolyte diffusion", "https://arxiv.org/abs/1910.10090"),
    "mace": ("Constructing MLIPs with minimum ab initio data", "https://www.nature.com/articles/s41524-026-02023-y"),
    "acs": ("Benchmarking universal MLIPs for alkali-ion battery kinetics", "https://doi.org/10.1021/acsmaterialslett.6c00134"),
    "ita": ("ITA research-area description", "https://www.pgfis.ita.br/en/post/fisica-atomica-e-molecular"),
    "efita": ("EFITA 2025 proceedings: neural-network potential for CaF2", "https://www.efita.ita.br/wp-content/uploads/2025/07/2025-07-31-Anais_do_XVIII_EFITA.pdf"),
    "pibic": ("ITA PIBIC 2025 result: superionic fluorides ML project", "https://paic.ita.br/documentos/editais/PIBIC-2025-Resultado.pdf"),
    "ice": ("Deep-Potential dataset for superionic water ice", "https://doi.org/10.5281/zenodo.6518808"),
    "rsnf": ("Russian Science Foundation project 23-73-01067", "https://rscf.ru/project/23-73-01067/"),
}

SEARCHES = [
    ("Web search / publisher indexes", '"Na2NbOF5" (MLIP OR MLMD OR "machine-learning force field" OR "neural network potential")', 4, 4, 0),
    ("Web search / publisher indexes", '"Na2NbOF5" (DeepMD OR "Deep Potential" OR MACE OR CHGNet OR M3GNet OR GAP)', 3, 3, 0),
    ("Web search / publisher indexes", '"Na2NbOF5" "machine learning" potential', 3, 3, 0),
    ("Web search / publisher indexes", '"Na2NbOF5" sodium diffusion AIMD conductivity', 7, 7, 3),
    ("Web search / publisher indexes", '"Na2NbOF5" molecular dynamics disorder conductivity', 6, 6, 2),
    ("Web search / publisher indexes", '"Na2NbOF5" (Haven ratio OR collective diffusion OR correlated hopping OR Nernst-Einstein)', 2, 2, 0),
    ("Web search / publisher indexes", '"Na2NbOF5" (defects OR vacancies OR supercell OR disorder)', 6, 6, 2),
    ("Web search / publisher indexes", '"Na2NbOF5" "classical potential"', 2, 2, 0),
    ("Crossref / DOI resolver", '10.1016/j.jpcs.2026.113889; 10.1007/s10008-025-06205-4; candidate contextual DOIs', 9, 9, 9),
    ("arXiv", '(sodium OR Na) AND (MLIP OR MLMD) AND (diffusion OR conductivity)', 8, 8, 3),
    ("Publisher indexes", '(sodium solid electrolyte OR superionic conductor) AND machine-learning interatomic potential AND transport', 10, 10, 5),
    ("ITA / institutional repositories", '"Filipe Matusalém" (machine learning potential OR superionic OR CaF2 OR water ice)', 8, 8, 4),
]

RECORDS = [
    ["R01", "A new sodium-ion conductor Na2NbOF5: Ab initio calculations and electrochemical testing", "Morkhova et al.", 2026, "Journal of Physics and Chemistry of Solids", "10.1016/j.jpcs.2026.113889", SOURCES["primary"][1], "Exact material", "Include", "Primary evidence", "Include"],
    ["R02", "A New Sodium-Ion Conductor Na2NbOF5: Ab Initio Calculations and Electrochemical Testing", "Morkhova et al.", 2026, "SSRN", "10.2139/ssrn.6481875", SOURCES["primary_preprint"][1], "Exact material", "Exclude", "Duplicate publication / preprint of R01", "Not assessed"],
    ["R03", "Hierarchical screening of ICSD to search for new high-conductive solid-state materials among mixed polyanionic oxohalides for sodium-ion batteries", "Orlova et al.", 2025, "Journal of Solid State Electrochemistry", "10.1007/s10008-025-06205-4", SOURCES["screening"][1], "Exact material", "Include", "Primary precursor study", "Include"],
    ["R04", "Na2NbOF5:Mn4+ red-emitting phosphor study", "Authors in source record", 2021, "Materials", "NR", SOURCES["phosphor"][1], "Exact material / different outcome", "Exclude", "Unrelated ion transport; no MLIP", "Not assessed"],
    ["R05", "Earlier synthesis/structure study of Na2NbOF5", "NR", "NR", "Structural literature", "NR", "NR", "Exact material / structure", "Exclude", "No molecular dynamics; unrelated transport", "Not assessed"],
    ["R06", "Computational Study of Na Diffusion and Conduction in P2- and O3-Na2x[NixTi1-x]O2 Materials with Machine-Learning Interatomic Potentials", "He; Chen; Lai", 2023, "SSRN", "10.2139/ssrn.4431516", SOURCES["he"][1], "Contextual sodium MLIP", "Include", "Method comparator", "Include"],
    ["R07", "Accurate Description of Ion Migration in Solid-State Ion Conductors from Machine-Learning Molecular Dynamics", "Miyagawa et al.", 2024, "arXiv", "10.48550/arXiv.2401.11244", SOURCES["miyagawa"][1], "Contextual MLMD", "Include", "Method comparator", "Include"],
    ["R08", "Probing Machine Learning Interatomic Potentials on Ion Transport Properties", "Aghoghovbia et al.", 2026, "Advanced Intelligent Discovery", "10.1002/aidi.70143", SOURCES["benchmark"][1], "Contextual benchmark", "Include", "Validation comparator", "Include"],
    ["R09", "Machine-learning-assisted discovery of lattice dynamics signatures of sodium superionic conductors", "Authors in source", 2025, "Materials Horizons", "10.1039/D5MH01176K", SOURCES["horizons"][1], "Contextual sodium MLIP", "Include", "Landscape comparator", "Include"],
    ["R10", "Simulating diffusion properties of solid-state electrolytes via a neural network potential: Performance and training scheme", "Authors in source", 2019, "arXiv", "10.48550/arXiv.1910.10090", SOURCES["deepmd"][1], "Contextual MLIP", "Include", "Method comparator", "Include"],
    ["R11", "Constructing machine learning interatomic potentials with minimum amount of ab initio data", "Authors in source", 2026, "npj Computational Materials", "10.1038/s41524-026-02023-y", SOURCES["mace"][1], "Contextual MLIP", "Include", "Validation comparator", "Include"],
    ["R12", "General review of artificial intelligence in batteries", "NR", 2024, "Review", "NR", "NR", "Broad review", "Include", "Potential context", "Exclude: review only / insufficient methodological relevance"],
    ["R13", "Desenvolvimento de um potencial interatômico baseado em redes neurais para o fluoreto superiônico CaF2", "Ramos; Matusalém", 2025, "XVIII EFITA proceedings", "NR", SOURCES["efita"][1], "Professor overlap", "Include", "Supplementary public evidence", "Include"],
    ["R14", "Plastic Deformation of Superionic Water Ices / Deep Potential dataset", "Matusalém et al.", 2022, "PNAS / Zenodo", "10.5281/zenodo.6518808", SOURCES["ice"][1], "Professor overlap", "Include", "Supplementary public evidence", "Include"],
]

DEDUP = [["D01", "R01", "R02", "Confirmed", "DOI differs by publication stage; exact normalized title and authors match", "Keep R01; retain R02 in audit log as preprint", DATE]]

EVIDENCE_HEADERS = ["Record ID", "Citation", "Year", "Material/system", "Mobile ion", "MLIP architecture", "Training-data source", "DFT method", "Training-set size", "Simulation cell size", "Trajectory length", "Temperature range", "Disorder/defects represented?", "Transport observable", "MSD", "Tracer diffusion", "Collective diffusion", "Ionic conductivity", "Haven ratio/correlation factor", "Activation energy", "Transport dimensionality", "Main finding", "Limitations", "Relevance", "Source URL"]
EVIDENCE = [
    ["R01", "Morkhova et al., JPCS 218, 113889", 2026, "Na2NbOF5", "Na+", "None (AIMD)", "N/A", "PBE; plane-wave cutoff 600 eV", "N/A", "Disorder-representing supercells; exact atom count awaiting supplied full text", "20 or 30 ps", "300 K", "Yes; O/F disorder represented", "AIMD diffusion and Nernst-Einstein conductivity", "Yes", "D≈7.1×10^-8 cm2 s^-1", "Not confirmed", "≈6.6×10^-3 S cm^-1", "Not confirmed", "NR", "2D", "Short AIMD predicts fast 2D Na transport", "Full text was not supplied; settings here are from user-provided paper notes plus accessible metadata. Experimental pellet is non-intrinsic.", "Primary material baseline", SOURCES["primary"][1]],
    ["R03", "Orlova et al., J Solid State Electrochem 29, 3043–3050", 2025, "Six mixed polyanionic oxohalides incl. Na2NbOF5", "Na+", "None", "N/A", "DFT reported; verify exact settings in full text", "N/A", "NR", "Kinetic Monte Carlo rather than MD trajectory", "Standard conditions and screening range", "Crystal structures screened", "BVSE barriers; kMC conductivity; vacancy formation", "No", "No", "No", "Predicted >10^-3 S cm^-1 for Na2NbOF5", "No", "Migration barriers reported (<0.5 eV in accessible summary)", "2D for Na2NbOF5", "Identified Na2NbOF5 as strongest screening candidate", "Screening-level approximations; not long-timescale atomistic MLMD", "Establishes prior discovery route, prevents novelty claim based on material alone", SOURCES["screening"][1]],
    ["R06", "He, Chen & Lai", 2023, "P2/O3 Na2x[NixTi1-x]O2", "Na+", "Neural-network MLIP", "DFT-labelled configurations", "NR in accessible record", "NR", "Large MLMD cells; exact values require full paper", "Long MLMD; exact values require full paper", "Temperature series", "Composition/configuration variation", "MSD, conductivity, residence/jumps, incoherent density correlations", "Yes", "Yes", "Yes/transport correlations", "Yes", "Yes; values below unity reported", "Yes", "Layered / anisotropic", "Shows correlated Na motion can invalidate simple independent-ion assumptions", "Preprint record; different electrode material", "Closest transport-analysis comparator; same methods, different physical system", SOURCES["he"][1]],
    ["R07", "Miyagawa et al.", 2024, "AgI; Na3SbS4; Li10GeP2S12", "Ag+/Na+/Li+", "On-the-fly VASP MLFF", "On-the-fly DFT", "PBE reported", "System-dependent", "System-dependent", "Longer than AIMD; system-dependent", "Multiple temperatures", "Disorder/defects represented where relevant", "Diffusion and concerted migration", "Yes", "Yes", "Material-dependent", "Material-dependent", "Material-dependent", "Yes", "Material-dependent", "MLMD reproduces ion migration across representative conductors", "No Na2NbOF5; transferability remains system-specific", "Supports feasibility but not exact-material novelty", SOURCES["miyagawa"][1]],
    ["R08", "Aghoghovbia et al.", 2026, "Li/Na superionic conductors", "Li+/Na+", "CHGNet; EquiformerV2; MatterSim; SevenNet; MACE", "Universal pretrained models", "Model-dependent", "N/A / pretrained", "Benchmark systems", "Transport trajectories", "Temperature-dependent", "Structure-dependent", "Ion transport benchmarking", "Yes", "Yes", "Varies", "Yes", "Varies", "Yes", "Varies", "Force accuracy alone does not guarantee accurate transport", "Universal models require transport-specific validation", "Key warning against unvalidated zero-shot MLMD", SOURCES["benchmark"][1]],
    ["R09", "Materials Horizons sodium-superionic survey", 2025, "921 sodium structures", "Na+", "Universal MLIP", "Pretrained plus database structures", "Model-dependent", "N/A / pretrained", "High-throughput cells", "Screening trajectories", "High-temperature / screening conditions", "Broad structural diversity", "Diffusivity and lattice dynamics descriptors", "Yes", "Yes", "No explicit exact-material result located", "Screening-level", "No exact-material result located", "Screening-level", "3D/varied", "Maps general signatures of Na superionic behavior", "Breadth rather than convergence/correlation depth", "Shows methodology is becoming routine; does not answer Na2NbOF5 question", SOURCES["horizons"][1]],
    ["R10", "Neural-network solid-electrolyte diffusion study", 2019, "Solid-state electrolytes", "Li+/Na+ depending system", "Deep Potential", "AIMD/DFT-labelled data", "System-dependent", "Reported in paper", "Larger than AIMD", "Long MLMD", "Temperature series", "System-dependent", "Diffusion", "Yes", "Yes", "NR", "Derived", "NR", "Yes", "System-dependent", "Demonstrates training and validation requirements for diffusion", "Different systems", "Method-development comparator", SOURCES["deepmd"][1]],
    ["R11", "npj Computational Materials MLIP data-efficiency study", 2026, "Ion-conducting benchmarks", "Mobile ions vary", "MACE / foundation-model route", "Ab initio labelled subsets", "Reported per benchmark", "Varies", "Varies", "Transport trajectories", "Varies", "Varies", "Diffusion", "Yes", "Yes", "NR", "Derived", "NR", "Yes", "Varies", "Data efficiency does not remove need to validate diffusion", "No Na2NbOF5 result", "Informs training design and error controls", SOURCES["mace"][1]],
    ["R13", "Ramos & Matusalém, EFITA abstract", 2025, "CaF2", "F-", "DeePMD-kit neural-network potential", "AIMD positions, forces and energies", "DFT/AIMD; exact functional NR in abstract", "NR", "96 atoms", "25 ps training trajectories", "1000–1600 K", "Superionic phase behavior", "Potential development for large systems/phase transitions", "NR", "NR", "NR", "NR", "NR", "NR", "Fluoride superionic", "Public evidence of MLIP work on CaF2", "Conference abstract only; not Na2NbOF5", "Adjacent methodological overlap", SOURCES["efita"][1]],
    ["R14", "Matusalém et al., superionic water ice", 2022, "Superionic water ice", "Protons/oxygen framework", "Deep Potential", "DFT-labelled data", "Reported in associated paper", "Public dataset", "Large MLMD systems", "Longer-scale deformation MD", "High-pressure/high-temperature", "Superionic phase", "Mechanical response and phase behavior", "NR", "NR", "NR", "NR", "NR", "NR", "Superionic phase", "MLIP enables deformation study beyond DFT scale", "Different chemistry and question", "Adjacent methodology, no direct material/problem overlap", SOURCES["ice"][1]],
]

GAPS = [
    ["A", "Trajectory-length convergence of Na diffusion", "SUPPORTED", "Primary study used 20/30 ps according to supplied notes; no exact-material convergence study located", "R06/R07", "Na2NbOF5 convergence remains material-specific", "Primary PDF needed to confirm all runs and block-analysis details"],
    ["B", "Finite-size effects and larger cells", "PARTIALLY SUPPORTED", "No cell-size convergence study located; prior work did use disorder-representing supercells", "R07/R08", "Question is convergence, not whether disorder was ignored", "Exact original cell sizes and replicas require full text"],
    ["C", "Sensitivity to O/F disorder realizations", "PARTIALLY SUPPORTED", "Disorder was represented, but systematic ensemble sampling was not located", "R07", "Multiple Na2NbOF5 O/F configurations are system-specific", "Full paper may contain more configuration sampling than accessible metadata shows"],
    ["D", "Collective/correlated hopping and Nernst-Einstein limits", "SUPPORTED", "No exact-material Haven/collective study located; contextual Na MLIP study demonstrates relevance", "R06", "Applies correlation analysis to Na2NbOF5 rather than importing conclusions", "Explicit treatment in primary paper must be confirmed from supplied PDF"],
    ["E", "Temperature-dependent diffusion and activation energy", "SUPPORTED", "Accessible exact-material AIMD evidence centers on 300 K; no long-trajectory temperature series located", "R06/R07", "Long MLMD can estimate a material-specific Arrhenius regime", "Confirm whether supplementary information contains additional temperatures"],
    ["F", "Defect/vacancy effects", "NOT SUPPORTED", "Prior screening addressed vacancy energetics and defect chemistry is not yet scoped well enough for a primary novelty claim", "R03/R04", "Could become a secondary sensitivity study only after defect states are justified", "Relevant intrinsic/extrinsic defect concentrations are unknown"],
]


def md_sources(keys):
    return "\n".join(f"- [{SOURCES[k][0]}]({SOURCES[k][1]})" for k in keys)


def write_md(name, text):
    (OUT / name).write_text(dedent(text).strip() + "\n", encoding="utf-8")


def csv_write(name, headers, rows):
    with (OUT / name).open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(color=WHITE, bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        max_len = max(len(str(ws.cell(row, col).value or "")) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(52, max(11, max_len + 2))
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 42
        for c in ws[row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if row % 2 == 0:
                c.fill = PatternFill("solid", fgColor=PALE)
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.style = "Hyperlink"


def workbook(name, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_sheet(ws)
    wb.save(OUT / name)


def fig_start(title, subtitle="", figsize=(11.7, 8.3)):
    fig, ax = plt.subplots(figsize=figsize)
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    ax.text(.04, .95, title, fontsize=20, weight="bold", color="#16324F", va="top")
    if subtitle:
        ax.text(.04, .905, subtitle, fontsize=10, color="#52616B", va="top")
    return fig, ax


def box(ax, x, y, w, h, title, body="", color="#2A7F83"):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.012,rounding_size=.012",facecolor="white",edgecolor=color,linewidth=1.8))
    ax.text(x+.018,y+h-.025,title,fontsize=11,weight="bold",color=color,va="top")
    ax.text(x+.018,y+h-.06,body,fontsize=8.5,color="#1F2933",va="top",wrap=True,linespacing=1.35)


def save_fig(fig, stem):
    fig.savefig(FIG / f"{stem}.png", dpi=320, bbox_inches="tight", facecolor="white")
    fig.savefig(FIG / f"{stem}.svg", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def visuals():
    fig, ax = fig_start("Focused scoping-review workflow", "Audit trail used for this real validation case")
    labels = [("1. Protocol","Falsification questions\nand eligibility"),("2. Search","12 exact query families\n6 source classes"),("3. Deduplicate","DOI → exact title →\nfuzzy title; log every action"),("4. Screen","14 raw records\n10 evidence sources"),("5. Extract","Material, MLIP, scale,\ndisorder and transport"),("6. Gap test","Six candidate gaps\nrated, not scored")]
    for i,(t,b) in enumerate(labels):
        x=.05+(i%3)*.31; y=.59-(i//3)*.28
        box(ax,x,y,.26,.16,t,b)
        if i%3<2: ax.annotate("",xy=(x+.305,y+.08),xytext=(x+.265,y+.08),arrowprops=dict(arrowstyle="->",color="#52616B",lw=1.5))
    ax.text(.05,.12,"Scientific integrity rule",fontsize=11,weight="bold",color="#A64040")
    ax.text(.05,.075,"Absence means “not located in the captured search”, never proof of non-existence. Full-paper claims remain pending until the supplied PDF is audited.",fontsize=9,color="#1F2933")
    save_fig(fig,"review_workflow")

    fig, ax = fig_start("PRISMA-style flow", "Focused validation corpus; counts are reconciled")
    boxes=[(.08,.68,"Records identified","n = 14"),(.08,.48,"After deduplication","n = 13"),(.08,.28,"Full texts assessed","n = 11"),(.08,.08,"Evidence sources included","n = 10")]
    for x,y,t,b in boxes: box(ax,x,y,.38,.12,t,b)
    for y in (.62,.42,.22): ax.annotate("",xy=(.27,y-.02),xytext=(.27,y+.05),arrowprops=dict(arrowstyle="->",color="#52616B"))
    box(ax,.58,.50,.34,.12,"Title/abstract exclusions","n = 2\nNo transport/MLIP relevance",color="#A64040")
    box(ax,.58,.28,.34,.12,"Full-text exclusion","n = 1\nReview only / insufficient relevance",color="#A64040")
    ax.annotate("",xy=(.58,.56),xytext=(.46,.56),arrowprops=dict(arrowstyle="->",color="#A64040")); ax.annotate("",xy=(.58,.34),xytext=(.46,.34),arrowprops=dict(arrowstyle="->",color="#A64040"))
    save_fig(fig,"prisma_flow")

    fig, ax = fig_start("Evidence map", "Exact-material depth versus methodological relevance")
    points=[("Primary Na2NbOF5",.9,.68,120,"#2A7F83"),("ICSD screening",.8,.42,90,"#2A7F83"),("Na layered MLIP",.25,.82,100,"#C58A19"),("Ion-conductor MLMD",.18,.7,90,"#C58A19"),("Universal MLIP benchmarks",.12,.54,110,"#C58A19"),("Prof. CaF2 MLIP",.2,.38,90,"#A64040"),("Superionic ice",.12,.27,90,"#A64040")]
    for label,x,y,s,c in points: ax.scatter(x,y,s=s,c=c); ax.text(x+.02,y,label,fontsize=8.5,va="center")
    ax.arrow(.08,.15,.8,0,width=.002,head_width=.018,color="#52616B"); ax.arrow(.08,.15,0,.68,width=.002,head_width=.018,color="#52616B")
    ax.text(.48,.09,"Specificity to Na2NbOF5 →",fontsize=10); ax.text(.025,.45,"Transport-analysis depth →",rotation=90,fontsize=10)
    save_fig(fig,"evidence_map")

    fig, ax = fig_start("Na2NbOF5 research timeline", "Publicly located milestones; not a claim of exhaustive historical coverage")
    ax.plot([.1,.9],[.48,.48],color="#52616B",lw=2)
    events=[(.15,"2021","Structure/luminescence\ncontext"),(.4,"2025","ICSD screening:\nBVSE + kMC"),(.68,"2026","AIMD + electrochemical\ntesting"),(.88,"Proposed","Validated long-timescale\nMLMD")]
    for x,year,lab in events:
        ax.scatter(x,.48,s=120,c="#2A7F83"); ax.text(x,.57,year,ha="center",weight="bold",color="#16324F"); ax.text(x,.4,lab,ha="center",va="top",fontsize=8.5)
    save_fig(fig,"research_timeline")

    fig, ax = fig_start("AIMD → MLIP → long-timescale MLMD", "Proposed validation chain; arrows do not imply automatic accuracy")
    stages=[(.05,"DFT/AIMD reference","Multiple O/F configurations\nenergies, forces, stresses"),(.285,"Train + validate MLIP","Held-out errors and structural/\ndynamical observables"),(.52,"Convergence ladder","Time blocks, cell sizes,\nconfiguration replicas"),(.755,"Transport analysis","Tracer + collective diffusion,\nHaven ratio, anisotropy, Ea")]
    for x,t,b in stages: box(ax,x,.42,.19,.2,t,b)
    for x in (.24,.475,.71): ax.annotate("",xy=(x+.035,.52),xytext=(x,.52),arrowprops=dict(arrowstyle="->",color="#52616B",lw=1.7))
    ax.text(.05,.25,"Required stop condition",weight="bold",color="#A64040"); ax.text(.05,.2,"Do not interpret long trajectories if the potential fails transport-relevant validation or extrapolates beyond its training domain.",fontsize=9)
    save_fig(fig,"aimd_to_mlip_workflow")

    fig, ax = fig_start("Candidate-gap matrix", "Evidence status after focused screening")
    status={"SUPPORTED":"#2A7F83","PARTIALLY SUPPORTED":"#C58A19","NOT SUPPORTED":"#A64040"}
    for i,(code,name,st,*_) in enumerate(GAPS):
        y=.78-i*.11; ax.text(.05,y,f"Gap {code}",weight="bold",fontsize=10); ax.text(.14,y,name,fontsize=9)
        ax.add_patch(FancyBboxPatch((.73,y-.022),.2,.05,boxstyle="round,pad=.006",facecolor=status[st],edgecolor="none")); ax.text(.83,y+.003,st,ha="center",va="center",fontsize=8,color="white",weight="bold")
    save_fig(fig,"candidate_gap_matrix")

    fig, ax = fig_start("Existing knowledge vs proposed contribution", "Boundary of defensible claims")
    box(ax,.06,.35,.39,.37,"Established", "• Na2NbOF5 was already selected by hierarchical screening\n• 2D Na pathways and short AIMD transport were reported\n• Disorder-representing supercells were used\n• MLIP transport workflows exist for other Na conductors", "#2A7F83")
    box(ax,.55,.35,.39,.37,"Proposed contribution", "• Demonstrate time and cell-size convergence\n• Sample multiple O/F disorder realizations\n• Separate tracer from collective charge transport\n• Determine temperature-dependent anisotropy and activation energy", "#C58A19")
    ax.text(.06,.22,"Not claimed",weight="bold",color="#A64040"); ax.text(.06,.17,"That disorder was ignored; that MLMD explains the experiment gap; or that a new material alone creates novelty.",fontsize=9)
    save_fig(fig,"knowledge_vs_contribution")


def main():
    search_rows=[]
    for i,(db,q,captured,screened,retained) in enumerate(SEARCHES,1):
        search_rows.append([f"S{i:02d}",db,q,DATE,"NR (interface did not expose total hit count)",captured,screened,retained,"Captured-hit counts are auditable, not database-total counts."])
    search_headers=["Search ID","Source/database","Exact search string","Search date","Reported total hits","Hits captured","Number screened","Number retained","Notes"]
    screen_headers=["Record ID","Title","Authors","Year","Venue","DOI","URL","Scope","Title/abstract decision","Reason / role","Full-text decision"]
    workbook("04_search_log.xlsx",[("Search Log",search_headers,search_rows),("Source Registry",["Key","Source","URL"],[[k,v[0],v[1]] for k,v in SOURCES.items()])])
    workbook("05_screening_log.xlsx",[("Records",screen_headers,RECORDS),("Deduplication",["Decision ID","Record A","Record B","Classification","Evidence","Resolution","Date"],DEDUP),("Reconciliation",["Measure","Count"],[["Raw records",14],["Duplicates",1],["Title/abstract screened",13],["Title/abstract excluded",2],["Full text assessed",11],["Full text excluded",1],["Included evidence sources",10]])])
    workbook("06_evidence_table.xlsx",[("Evidence",EVIDENCE_HEADERS,EVIDENCE),("Gap Matrix",["Gap","Candidate gap","Status","Evidence","Closest comparator","Distinctness","Remaining uncertainty"],GAPS),("Field Guide",["Term","Meaning"],[["NR","Not reported or not verified in accessible source"],["Not confirmed","Requires full-paper audit"],["Contextual","Informs method but is not evidence that Na2NbOF5 has been studied"]])])
    csv_write("search_log.csv",search_headers,search_rows); csv_write("screening_log.csv",screen_headers,RECORDS); csv_write("evidence_table.csv",EVIDENCE_HEADERS,EVIDENCE); csv_write("gap_matrix.csv",["Gap","Candidate gap","Status","Evidence","Closest comparator","Distinctness","Remaining uncertainty"],GAPS)

    write_md("01_executive_summary.md", f'''# Executive summary

**Review date:** {DATE}  
**Review type:** focused scoping review and gap analysis  
**Decision standard:** attempt to falsify the proposal.

## Bottom line

The captured search did **not locate a dedicated MLIP/MLMD study of Na2NbOF5**, a long-trajectory convergence study, a classical potential, or an exact-material analysis of Haven ratio/collective conductivity. This is an absence-of-located-evidence statement, not proof of absence. The proposal is defensible only if it is framed around **validated convergence and correlation analysis**, not around the material merely being new.

The strongest surviving gap is: **whether Na2NbOF5 transport estimates are stable with respect to trajectory length, finite cell size and O/F disorder realization, and whether collective charge transport differs materially from tracer/Nernst–Einstein estimates across temperature.**

## What is already established

- Hierarchical ICSD screening, BVSE and kinetic Monte Carlo already selected Na2NbOF5 and predicted favorable two-dimensional Na transport.
- The 2026 primary study reported short AIMD transport at 300 K. User-supplied notes give 20/30 ps, 1 fs, NVT, Nosé–Hoover, PBE, 600 eV, D ≈ 7.1 × 10^-8 cm² s^-1 and conductivity ≈ 6.6 × 10^-3 S cm^-1.
- The prior work **did represent structural disorder** using supercells. This review does not claim otherwise.
- MLIP/MLMD transport analysis is established in other Na-ion and superionic systems; it is an enabling method, not novelty by itself.

## Important limitation

The requested primary-paper full text was not present in the workspace and publisher PDF access was unavailable. Therefore the primary-paper audit labels values from the user's notes separately and leaves unverified items unresolved. No statement is attributed to the authors unless supported by an accessible record.

## Overlap with Prof. Filipe Matusalém

Classification: **B — adjacent overlap.** Public evidence shows ML interatomic potentials for superionic CaF2/related fluorides and superionic water ice, but no direct Na2NbOF5 project or the same Na-transport convergence/correlation question was located. Unpublished plans can only be resolved by asking the supervisor.

## Gap status

- Supported: A trajectory length; D correlations; E temperature-dependent long trajectories.
- Partially supported: B finite size; C disorder realizations.
- Not supported as a primary gap: F defects/vacancies.

## Sources
{md_sources(["primary","screening","he","miyagawa","benchmark","efita","ice"])}''')

    write_md("02_review_protocol.md", '''# Review protocol

## Objective

Determine whether a Na2NbOF5 MLIP/MLMD proposal remains scientifically distinct after searching exact-material work, contextual Na-ion MLIP literature and the supervisor's public research.

## Eligibility

**Include:** exact-material synthesis, screening, atomistic transport or disorder studies; MLIP/MLMD studies that directly analyze solid-state ion transport; public supervisor outputs that establish material/problem overlap.  
**Exclude:** unrelated properties without transport relevance; generic AI/battery reviews; records with no molecular dynamics or methodological relevance; duplicate publication; conference abstracts as primary scientific evidence (they may be supplementary overlap evidence).

## Outcomes

Exact-material MLIP use; simulation duration and size; disorder sampling; tracer and collective diffusion; Nernst–Einstein use; Haven/correlation analysis; activation energy; classical-potential availability; competing groups; supervisor overlap.

## Deduplication

DOI was the primary key, followed by normalized exact title and then fuzzy title review. Candidate duplicates were retained in the log. R02 was retained as an audit record but excluded as the preprint/publication version of R01.

## Screening

Title/abstract decisions were Include, Exclude or Maybe. Every exclusion has a reason. Full-text exclusions use protocol categories. Contextual papers were retained only when they tested transport-relevant MLIP capabilities or limitations.

## Synthesis

Candidate gaps were classified SUPPORTED, PARTIALLY SUPPORTED or NOT SUPPORTED. No numerical novelty score was used. Negative search results are phrased as “not located”.

## Deviations and limitations

- Search interfaces exposed a ranked result set but not reproducible global hit counts. The log therefore records `NR` for total hits and separately reports captured/screened records.
- The primary paper PDF was not supplied. Its audit is provisional and provenance-labelled.
- Citation chasing and subscription database exports were not available; this is a focused validation, not a claim of exhaustive systematic coverage.
''')

    write_md("03_search_strategy.md", f'''# Search strategy

## Sources searched

Web/publisher indexes, Crossref/DOI resolver, arXiv, publisher journal pages, SSRN, PubMed, ITA institutional pages/proceedings and the Russian Science Foundation project page. Search date: **{DATE}**.

## Query logic

Exact-material queries combined `Na2NbOF5` with MLIP synonyms, named architectures, transport observables, disorder/defects and classical-potential terms. Contextual queries combined sodium solid electrolytes/superionic conductors with MLIP and transport. The professor search combined the name with ML potentials and known systems.

The exact strings, captured counts, screened counts and retained counts are in `04_search_log.xlsx`. Total result counts are `NR` where the interface did not expose them; inventing a count would be less reproducible than reporting the limitation.

## Exact-material conclusion

No captured result established a dedicated Na2NbOF5 MLIP, long-timescale MLMD, Haven-ratio study, convergence study or classical potential. The original synthesis/screening group remains an obvious competing group, and its ongoing work should be checked directly before proposal submission.

## Key source registry
{md_sources(["primary","primary_preprint","screening","phosphor","he","miyagawa","benchmark","horizons","deepmd","mace","rsnf"])}''')

    write_md("07_gap_analysis.md", f'''# Scientific gap analysis

## Surviving gap

**Whether Na2NbOF5 transport predictions converge with trajectory length and cell size across multiple O/F disorder realizations, and whether collective charge transport differs from tracer/Nernst–Einstein transport across temperature.**

This combines the supported gaps without claiming that prior AIMD ignored disorder or that MLMD must reconcile calculation and experiment.

## Candidate assessments

''' + "\n\n".join(f'''### Gap {g[0]} — {g[1]}

**Status:** {g[2]}  
**Evidence:** {g[3]}  
**Closest competing paper:** {g[4]}  
**Why distinct:** {g[5]}  
**Remaining uncertainty:** {g[6]}''' for g in GAPS) + f'''

## Claims to avoid

- “Previous AIMD ignored disorder.” It did not.
- “MLMD will explain the calculation–experiment discrepancy.” The reported experimental sample is not an intrinsic bulk benchmark.
- “The work is novel because Na2NbOF5 is new.” Prior screening and AIMD already exist.
- “No one is working on this.” Public search cannot resolve unpublished work.

## Sources
{md_sources(["primary","screening","he","miyagawa","benchmark","horizons"])}''')

    write_md("08_professor_overlap_analysis.md", f'''# Professor overlap analysis

## Classification

**B. Adjacent overlap — same broad methodology, different system/question.**

## Evidence

- ITA publicly describes Prof. Filipe Matusalém's research as DFT, molecular dynamics and machine learning applied to atomistic systems.
- An EFITA 2025 abstract with Davi Ramos describes a DeePMD-kit potential for superionic CaF2 trained from 96-atom AIMD cells at 1000–1600 K over 25 ps.
- An ITA PIBIC listing covers ML atomistic simulation of CaF2, PbF2 and CdF2.
- Public work/datasets document Deep-Potential simulations of superionic water ice.

## Interpretation

The overlap is methodological and thematic: MLIPs, superionic transport and extending DFT-scale simulations. The located public record does **not** show Na2NbOF5, Na-ion transport in an oxyfluoride, O/F configuration ensembles, or the same tracer-versus-collective transport question.

## Limitation and action

No public search can exclude an unpublished group project. Before submission, ask whether the group already has Na2NbOF5 data, a trained potential, or a student assigned to the system. This is a project-positioning question, not evidence that the proposal is duplicate.

## Sources
{md_sources(["ita","efita","pibic","ice"])}''')

    write_md("09_primary_paper_audit.md", f'''# Na2NbOF5 primary-paper audit

## Audit status

**PROVISIONAL — full paper not supplied in the workspace.** Publisher metadata and search snippets were accessible; the numerical simulation details below are explicitly from the user's paper notes. They must be checked line-by-line against the provided PDF before this audit can be called complete.

## FACT FROM ACCESSIBLE PAPER RECORD

- Title: *A new sodium-ion conductor Na2NbOF5: Ab initio calculations and electrochemical testing*.
- DOI: 10.1016/j.jpcs.2026.113889; Journal of Physics and Chemistry of Solids 218, 113889 (2026).
- The work follows hierarchical ICSD screening using geometrical/topological analysis, bond-valence site-energy ideas, kinetic Monte Carlo and quantum-chemical calculations.
- Accessible summaries identify two-dimensional Na migration and barriers below 0.5 eV.

## FACT FROM USER-PROVIDED FULL-PAPER NOTES — REQUIRES PDF VERIFICATION

- AIMD length: 20 or 30 ps; timestep: 1 fs; ensemble: NVT; thermostat: Nosé–Hoover; temperature: 300 K.
- DFT: PBE; plane-wave cutoff: 600 eV.
- Disorder: supercells were prepared to represent structural O/F disorder.
- Reported D: approximately 7.1 × 10^-8 cm² s^-1.
- Reported conductivity: approximately 6.6 × 10^-3 S cm^-1.
- Experimental sample: about 12.3% NaF; measured conductivity approximately 1.03 × 10^-6 S cm^-1.
- The experimental number was not presented as intrinsic bulk conductivity because the sample was unsintered, contacts were poor, grain boundaries contributed and NaF was present.

## ITEMS NOT VERIFIED WITHOUT FULL TEXT

- Exact reason for material selection in the final paper versus the prior screening study.
- Exact migration-path labels and numerical barrier for each path.
- Exact cell sizes, O/F occupation construction and number of configurations.
- MSD fitting window, dimensional prefactor, uncertainty/block averaging and derivation of D.
- Exact Nernst–Einstein equation, carrier concentration and charge used for conductivity.
- Whether distinct/self and collective conductivity, Haven ratio or an explicit correlation factor were calculated.
- Authors' stated future work.

## OUR INTERPRETATION

Twenty-to-thirty-picosecond trajectories can be too short to demonstrate diffusive-regime convergence, especially in disordered solids. That motivates a convergence study; it does not invalidate the paper. Multiple disorder realizations may reveal variance not visible in a single representation, but the paper must not be described as ignoring disorder. The experimental discrepancy is not a clean MLMD validation target because the sample limitations confound intrinsic transport.

## Sources
{md_sources(["primary","primary_preprint","screening"])}''')

    write_md("10_final_proposal_concept.md", '''# Final proposal concept

## Working title

**Disorder- and Correlation-Resolved Sodium Transport in Na2NbOF5 from Validated Machine-Learning Molecular Dynamics**

## Refined research question

**How do trajectory length, finite cell size, O/F configurational realization and collective Na-ion correlations affect tracer and collective transport estimates in Na2NbOF5 across temperature?**

## Three objectives

1. Develop and validate an MLIP against held-out DFT energies, forces and stresses and against short-AIMD structural/dynamical observables for multiple physically justified O/F configurations.
2. Quantify convergence of Na tracer diffusion, two-dimensional anisotropy and activation energy with trajectory length, cell size, temperature and disorder realization.
3. Compute collective charge transport, Haven/correlation factors and hopping statistics, and compare these with simple Nernst–Einstein estimates without treating the non-intrinsic pellet conductivity as a direct bulk benchmark.

## What MLIP specifically enables

Longer trajectories, larger cells, replicated disorder configurations and temperature sampling that are impractical with direct AIMD. It does not guarantee accuracy; transport-specific validation and uncertainty checks are part of the work.

## Supervisor-ready paragraph

The published work already establishes Na2NbOF5 as a candidate two-dimensional Na conductor and includes disorder-representing short AIMD, so I would not position the proposal as discovering the material or correcting an omission of disorder. The focused literature search did not locate a dedicated Na2NbOF5 MLIP/MLMD study, systematic time/cell/configuration convergence, or collective-versus-tracer transport analysis. I therefore propose to develop a validated MLIP and use it to test whether the reported transport is robust across longer trajectories, larger cells, multiple O/F configurations and temperature, including Haven/correlation analysis. This is methodologically adjacent to your MLIP work on superionic systems, but the located public record indicates a different material and physical question. Before finalizing, I would like to confirm whether the group already has unpublished Na2NbOF5 work or a preferred potential framework.

## Evidence still needed

The full primary PDF and supplementary information; a subscription-database/citation-chaining update; confirmation of the original group's current projects; supervisor confirmation of unpublished overlap; and a pre-registered MLIP validation threshold before production MLMD.
''')

    write_md("11_case_study_full.md", f'''# SciReview Chem — Real-World Validation Case Study

## From literature search to research-gap identification for MLIP simulation of Na2NbOF5

### Case purpose

This case tested whether SciReview Chem could support a real, adversarial proposal screen rather than produce a synthetic portfolio artifact. The target was a graduate proposal on long-timescale Na transport in disordered Na2NbOF5. The review attempted to falsify novelty across exact-material literature, contextual MLIP transport work and the prospective supervisor's public research.

### Protocol and search

Twelve documented query families were run across six source classes. Fourteen unique audit records were assembled; one preprint/journal duplicate was logged, two records were excluded at title/abstract, eleven were assessed for evidence relevance, one generic review was excluded at full text and ten sources were retained. Because the search interface did not expose reproducible database-total counts, the log reports captured and screened counts rather than invented global hit counts.

![Review workflow](visual_assets/review_workflow.png)

![PRISMA-style flow](visual_assets/prisma_flow.png)

### Evidence result

The exact-material literature located comprises a hierarchical screening study and the subsequent AIMD/electrochemical paper. The public record did not reveal a dedicated Na2NbOF5 MLIP, long-time convergence analysis, explicit Haven/collective treatment, or classical potential. Contextual studies show that MLMD, correlation analysis and universal MLIP screening are established methods. Therefore the method is not the contribution; the Na2NbOF5-specific convergence and correlation question may be.

![Evidence map](visual_assets/evidence_map.png)

### Primary-paper integrity boundary

The user specified exact AIMD and experimental details, but no PDF was present. SciReview Chem therefore records those data as user-provided notes pending source verification. The audit never claims that the paper ignored disorder and never treats the unsintered, impure pellet conductivity as intrinsic bulk validation.

### Gap test

Three candidate gaps were supported, two partly supported, and one not supported. The surviving question integrates trajectory/cell/configuration convergence with collective transport. Defects/vacancies were rejected as a primary gap because the accessible evidence does not yet justify a distinct, bounded claim.

![Candidate gap matrix](visual_assets/candidate_gap_matrix.png)

![Existing knowledge and contribution](visual_assets/knowledge_vs_contribution.png)

### Supervisor overlap

Public records show Prof. Matusalém using ML potentials for superionic CaF2/related fluorides and water ice. This is adjacent methodological overlap, not the same material/problem. Direct overlap was not located; unpublished overlap remains a question for the supervisor.

### App validation findings

The real case revealed that SciReview Chem's search log lacked screened/retained fields, title-stage exclusions did not require a reason, and extraction templates omitted MLIP transport variables. The application was updated to capture those counts, require exclusion reasons, and add an MLIP ion-transport template covering training data, simulation scale, disorder, tracer/collective diffusion, conductivity, Haven ratio and activation energy. Regression tests were added.

### Limitations

This is a focused scoping review based on captured web/publisher records, not an exhaustive subscription-database search. Total hit counts were unavailable, the full primary paper was not supplied, and public sources cannot exclude unpublished competing work. These limitations narrow the claims rather than being filled with inference.

### Final concept

The defensible project is a validated MLMD convergence and transport-correlation study, not a generic application of MLIP to a new material. The workflow is AIMD reference → validated potential → convergence ladder → tracer and collective transport analysis.

![AIMD to MLIP workflow](visual_assets/aimd_to_mlip_workflow.png)

### Source registry
{md_sources(list(SOURCES))}''')

    payload={"metadata":{"title":"SciReview Chem real-world validation: Na2NbOF5","review_date":DATE,"review_type":"focused scoping review / gap analysis","primary_pdf_supplied":False},"counts":{"source_classes":6,"search_families":12,"raw_records":14,"duplicates":1,"title_abstract_screened":13,"title_abstract_excluded":2,"full_text_assessed":11,"full_text_excluded":1,"included_evidence_sources":10,"supported_gaps":3,"partially_supported_gaps":2,"not_supported_gaps":1},"searches":[dict(zip(search_headers,row)) for row in search_rows],"records":[dict(zip(screen_headers,row)) for row in RECORDS],"duplicates":[dict(zip(["decision_id","record_a","record_b","classification","evidence","resolution","date"],row)) for row in DEDUP],"evidence":[dict(zip(EVIDENCE_HEADERS,row)) for row in EVIDENCE],"gaps":[dict(zip(["gap","candidate_gap","status","evidence","closest_comparator","distinctness","remaining_uncertainty"],row)) for row in GAPS],"surviving_gap":"Whether Na2NbOF5 transport predictions converge with trajectory length and cell size across multiple O/F disorder realizations, and whether collective charge transport differs from tracer/Nernst–Einstein transport across temperature.","professor_overlap":{"classification":"B — Adjacent overlap","direct_overlap_located":False,"limitation":"Unpublished work cannot be assessed from public sources."},"app_validation":{"bugs_discovered":3,"bugs_fixed":3,"fixes":["search screened/retained counts","mandatory title-exclusion reasons","MLIP ion-transport extraction template"]},"limitations":["Primary PDF not supplied","Search interface did not expose total hit counts","Focused public-web search is not exhaustive","Unpublished competing work cannot be ruled out"],"sources":{k:{"label":v[0],"url":v[1]} for k,v in SOURCES.items()}}
    (OUT/"12_case_study_data.json").write_text(json.dumps(payload,indent=2,ensure_ascii=False),encoding="utf-8")
    visuals()

    # Structural workbook verification.
    for name, expected in {"04_search_log.xlsx":{"Search Log","Source Registry"},"05_screening_log.xlsx":{"Records","Deduplication","Reconciliation"},"06_evidence_table.xlsx":{"Evidence","Gap Matrix","Field Guide"}}.items():
        wb=load_workbook(OUT/name,read_only=False,data_only=False)
        assert expected.issubset(wb.sheetnames)
        for ws in wb.worksheets:
            assert ws.freeze_panes == "A2"
            assert ws.auto_filter.ref
            assert ws.max_row >= 2 and ws.max_column >= 2
    print(json.dumps(payload["counts"],indent=2))


if __name__ == "__main__":
    main()
