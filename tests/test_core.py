from pathlib import Path

import pytest

from scireview.deduplication import apply_resolution, find_duplicate_candidates, normalize_doi, normalize_title
from scireview.importers import import_references
from scireview.extraction import ADSORPTION_FIELDS, MLIP_TRANSPORT_FIELDS, included_records, missing_fields, upsert_extraction
from scireview.evidence import build_evidence_table, filter_evidence_rows
from scireview.models import BibliographicRecord, ReviewProject, ScreeningDecision
from scireview.persistence import load_project, project_from_json, project_to_json, save_project
from scireview.prisma import derive_prisma_counts, export_prisma_figure
from scireview.screening import eligible_for_full_text, record_decision, screening_summary
from scireview.quality import upsert_quality_assessment
from scireview.synthesis import build_synthesis_summary, frequency_table
from scireview.reporting import export_project_workbook, export_workflow_pdf
from scireview.demo import DEMO_WARNING, build_demo_project
from scireview.utils import lines_to_list


def test_normalization():
    assert normalize_doi("https://doi.org/10.1000/ABC.") == "10.1000/abc"
    assert normalize_title("Fe–Oxide: Adsorption!") == "fe oxide adsorption"


def test_exact_doi_and_fuzzy_candidates_are_not_removed():
    records = [
        BibliographicRecord(title="A study of adsorption", doi="10.1/ABC"),
        BibliographicRecord(title="A study of adsorption", doi="https://doi.org/10.1/abc"),
        BibliographicRecord(title="A study of adsorption behaviour"),
    ]
    candidates = find_duplicate_candidates(records, fuzzy_threshold=75)
    assert candidates[0].classification == "confirmed"
    assert "doi" in candidates[0].matching_fields
    assert all(record.active for record in records)
    assert all(candidate.resolution is None for candidate in candidates)


def test_full_text_exclusion_requires_reason():
    decision = ScreeningDecision(record_id="x", stage="full_text", decision="exclude")
    with pytest.raises(ValueError, match="requires an explicit reason"):
        decision.validate()


def test_title_abstract_exclusion_requires_reason():
    decision = ScreeningDecision(record_id="x", stage="title_abstract", decision="exclude")
    with pytest.raises(ValueError, match="requires an explicit reason"):
        decision.validate()


def test_mlip_transport_template_covers_correlation_and_convergence_fields():
    keys = {field.key for field in MLIP_TRANSPORT_FIELDS}
    assert {"simulation_cell_size", "trajectory_length", "collective_diffusion", "haven_ratio", "disorder_defects"}.issubset(keys)


def test_project_round_trip(tmp_path: Path):
    project = ReviewProject(records=[BibliographicRecord(title="Preserved title", raw_metadata={"TI": "Raw"})])
    target = save_project(project, tmp_path / "review.json")
    loaded = load_project(target)
    assert loaded.records[0].title == "Preserved title"
    assert loaded.records[0].raw_metadata == {"TI": "Raw"}


def test_prisma_uses_latest_decisions_and_reports_pending():
    first = BibliographicRecord(title="Included")
    second = BibliographicRecord(title="Excluded")
    project = ReviewProject(records=[first, second])
    project.add_screening_decision(ScreeningDecision(record_id=first.id, decision="uncertain", decided_at="2026-01-01T00:00:00+00:00"))
    project.add_screening_decision(ScreeningDecision(record_id=first.id, decision="include", decided_at="2026-01-02T00:00:00+00:00"))
    project.add_screening_decision(ScreeningDecision(record_id=second.id, decision="exclude", reason="Wrong system"))
    project.add_screening_decision(ScreeningDecision(record_id=first.id, stage="full_text", decision="include"))
    counts = derive_prisma_counts(project)
    assert counts.records_screened == 2
    assert counts.records_excluded == 1
    assert counts.reports_sought == 1
    assert counts.studies_included == 1
    assert counts.reconciled


def test_empty_prisma_flow_is_not_presented_as_complete():
    counts = derive_prisma_counts(ReviewProject())
    assert not counts.reconciled
    assert "No records have been imported" in counts.warnings[0]


def test_lines_to_list_preserves_wording_and_order():
    assert lines_to_list("  First criterion  \n\nSecond criterion") == ["First criterion", "Second criterion"]


def test_csv_import_preserves_raw_metadata():
    payload = b'Title,Authors,Year,DOI,Custom\n"Imported study","Able, A.; Baker, B.",2024,10.1/demo,untouched\n'
    result = import_references(payload, "records.csv", "Example database")
    assert result.records[0].title == "Imported study"
    assert result.records[0].doi == "10.1/demo"
    assert result.records[0].raw_metadata["Custom"] == "untouched"
    assert result.records[0].source_database == "Example database"


def test_ris_import_handles_repeated_authors():
    payload = b"TY  - JOUR\nTI  - Transparent adsorption study\nAU  - Able, Ada\nAU  - Baker, Ben\nPY  - 2023\nDO  - 10.2/ris\nER  -\n"
    result = import_references(payload, "records.ris")
    assert result.records[0].authors == ["Able, Ada", "Baker, Ben"]
    assert result.records[0].year == 2023
    assert result.records[0].raw_metadata["TY"] == "JOUR"


def test_bibtex_import_maps_fields_without_inventing_values():
    payload = b'@article{demo, title={A {Nested} Title}, author={Able, Ada and Baker, Ben}, year={2022}, journal={Demo Journal}}'
    result = import_references(payload, "records.bib")
    record = result.records[0]
    assert record.title == "A {Nested} Title"
    assert record.authors == ["Able, Ada", "Baker, Ben"]
    assert record.doi is None
    assert record.raw_metadata["ID"] == "demo"


def test_duplicate_resolution_requires_human_action_and_preserves_records():
    left = BibliographicRecord(title="Same title", doi="10.1/x", journal=None, raw_metadata={"left": True})
    right = BibliographicRecord(title="Same title", doi="10.1/x", journal="Journal", raw_metadata={"right": True})
    project = ReviewProject(records=[left, right])
    project.duplicate_candidates = find_duplicate_candidates(project.records)
    assert left.active and right.active
    apply_resolution(project, project.duplicate_candidates[0].id, "merge", reviewer="AB")
    assert len(project.records) == 2
    assert left.active and not right.active
    assert left.journal == "Journal"
    assert left.raw_metadata == {"left": True}
    assert right.raw_metadata == {"right": True}
    assert project.duplicate_candidates[0].reviewer == "AB"


def test_screening_history_is_append_only_and_latest_decision_wins():
    record = BibliographicRecord(title="Review me")
    project = ReviewProject(records=[record])
    record_decision(project, record.id, "title_abstract", "uncertain", notes="First pass")
    record_decision(project, record.id, "title_abstract", "include", notes="Resolved")
    assert len(project.screening_decisions) == 2
    assert project.latest_decision(record.id, "title_abstract").decision == "include"
    summary = screening_summary(project, "title_abstract")
    assert summary.included == 1 and summary.uncertain == 0 and summary.pending == 0


def test_full_text_eligibility_and_exclusion_reason_validation():
    included = BibliographicRecord(title="Eligible")
    excluded = BibliographicRecord(title="Not eligible")
    project = ReviewProject(records=[included, excluded])
    record_decision(project, included.id, "title_abstract", "include")
    record_decision(project, excluded.id, "title_abstract", "exclude", "Wrong system")
    assert eligible_for_full_text(project) == {included.id}
    with pytest.raises(ValueError, match="requires an explicit reason"):
        record_decision(project, included.id, "full_text", "exclude")
    with pytest.raises(ValueError, match="not eligible"):
        record_decision(project, excluded.id, "full_text", "include")
    record_decision(project, included.id, "full_text", "exclude", "No relevant outcome")
    assert screening_summary(project, "full_text").excluded == 1


def _included_project():
    record = BibliographicRecord(title="Included study")
    project = ReviewProject(records=[record])
    record_decision(project, record.id, "title_abstract", "include")
    record_decision(project, record.id, "full_text", "include")
    return project, record


def test_extraction_is_limited_to_included_studies_and_missing_is_not_zero():
    project, record = _included_project()
    assert included_records(project) == [record]
    entry = upsert_extraction(project, record.id, "adsorption", {"ph": "", "adsorption_capacity": "0"})
    assert entry.values["ph"] is None
    assert entry.values["adsorption_capacity"] == "0"
    assert "ph" in missing_fields(entry, ADSORPTION_FIELDS)
    assert "adsorption_capacity" not in missing_fields(entry, ADSORPTION_FIELDS)


def test_quality_assessment_stores_domains_without_aggregate_score():
    project, record = _included_project()
    assessment = upsert_quality_assessment(
        project, record.id, "Context-specific",
        {"Method description": ("some_concern", "Important detail absent")}, reviewer="AB",
    )
    assert assessment.domains[0].rating == "some_concern"
    assert not hasattr(assessment, "score")
    with pytest.raises(ValueError, match="Invalid quality rating"):
        upsert_quality_assessment(project, record.id, "Bad", {"Domain": ("excellent", "")})


def test_evidence_table_preserves_zero_and_marks_missing():
    project, record = _included_project()
    upsert_extraction(project, record.id, "adsorption", {"ph": None, "adsorption_capacity": "0", "kinetic_models": "Pseudo-first; Pseudo-second"})
    table = build_evidence_table(project)
    assert table.rows[0]["ph"] == "NR"
    assert table.rows[0]["adsorption_capacity"] == "0"
    assert filter_evidence_rows(table, "included study") == table.rows
    assert filter_evidence_rows(table, filters={"adsorption_capacity": {"1"}}) == []


def test_synthesis_uses_only_reported_values():
    first = BibliographicRecord(title="One", year=2020)
    second = BibliographicRecord(title="Two", year=2020)
    project = ReviewProject(records=[first, second])
    for record in project.records:
        record_decision(project, record.id, "title_abstract", "include")
        record_decision(project, record.id, "full_text", "include")
    upsert_extraction(project, first.id, "adsorption", {"equilibrium_models": "Langmuir; Freundlich", "bet_surface_area": "NR"})
    upsert_extraction(project, second.id, "adsorption", {"equilibrium_models": "Langmuir", "bet_surface_area": "0"})
    table = build_evidence_table(project)
    assert frequency_table(table, "equilibrium_models", split_multi=True) == {"Langmuir": 2, "Freundlich": 1}
    summary = build_synthesis_summary(table)
    assert summary.year_counts == {"2020": 2}
    assert summary.missing_counts["bet_surface_area"] == 1


def test_prisma_reason_counts_use_current_decisions():
    included = BibliographicRecord(title="Included")
    excluded = BibliographicRecord(title="Excluded")
    project = ReviewProject(records=[included, excluded])
    record_decision(project, included.id, "title_abstract", "include")
    record_decision(project, excluded.id, "title_abstract", "exclude", "Wrong outcome")
    record_decision(project, included.id, "full_text", "exclude", "No relevant outcome")
    counts = derive_prisma_counts(project)
    assert counts.title_exclusion_reasons == {"Wrong outcome": 1}
    assert counts.full_text_exclusion_reasons == {"No relevant outcome": 1}
    assert counts.reconciled


@pytest.mark.parametrize("format,signature", [("png", b"\x89PNG"), ("svg", b"<?xml"), ("pdf", b"%PDF")])
def test_prisma_figure_exports(format, signature):
    project, _ = _included_project()
    payload = export_prisma_figure(derive_prisma_counts(project), format, "Test review")
    assert payload.startswith(signature)
    assert len(payload) > 1000


def test_project_json_download_round_trip():
    project, record = _included_project()
    loaded = project_from_json(project_to_json(project))
    assert loaded.records[0].id == record.id
    assert loaded.latest_decision(record.id, "full_text").decision == "include"


def test_excel_export_contains_required_sheets_and_raw_metadata():
    from io import BytesIO
    from openpyxl import load_workbook

    project, record = _included_project()
    record.raw_metadata = {"original_title": "Untouched source value"}
    workbook = load_workbook(BytesIO(export_project_workbook(project)), read_only=True, data_only=False)
    required = {"Protocol", "Search Strategies", "Imported Records", "Deduplication Log", "Screening Decisions",
                "Full Text Decisions", "Included Studies", "Data Extraction", "Quality Assessment", "Evidence Table",
                "PRISMA Counts", "Synthesis Summary", "Metadata"}
    assert required.issubset(workbook.sheetnames)
    assert workbook["Imported Records"]["A2"].value == record.id
    assert "Untouched source value" in workbook["Imported Records"]["O2"].value


def test_pdf_workflow_report_is_readable():
    from io import BytesIO
    from pypdf import PdfReader

    project, _ = _included_project()
    project.protocol.title = "Transparent review"
    payload = export_workflow_pdf(project)
    assert payload.startswith(b"%PDF")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages)
    assert "Transparent review" in text
    assert "Included-study evidence" in text


def test_demo_project_is_complete_and_unmistakably_synthetic():
    demo = build_demo_project()
    assert demo.metadata.demonstration
    assert len(demo.records) == 40
    assert all(record.synthetic and DEMO_WARNING in record.raw_metadata["warning"] for record in demo.records)
    counts = derive_prisma_counts(demo)
    assert counts.duplicates_removed == 4
    assert counts.records_excluded == 10
    assert counts.full_texts_excluded == 6
    assert counts.studies_included == 20
    assert counts.reconciled
    assert len(demo.extraction_entries) == 20
    assert len(demo.quality_assessments) == 20
