from __future__ import annotations

from io import BytesIO
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scireview.evidence import build_evidence_table
from scireview.models import ReviewProject
from scireview.prisma import derive_prisma_counts
from scireview.screening import current_decisions


NAVY = "0B3558"
TEAL = "159E9C"
PALE = "EAF4F4"
WHITE = "FFFFFF"


def _scalar(value: Any) -> Any:
    if value is None:
        return "NR"
    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(item) for item in value) or "NR"
    if isinstance(value, dict):
        return "; ".join(f"{key}: {item}" for key, item in value.items()) or "NR"
    return value


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_scalar(value) for value in row])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, column in enumerate(sheet.columns, start=1):
        width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        for cell in column:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_index in range(2, sheet.max_row + 1):
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor=PALE)


def export_project_workbook(project: ReviewProject) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    protocol = project.protocol
    _sheet(workbook, "Protocol", ["Field", "Value"], [[key.replace("_", " ").title(), value] for key, value in vars(protocol).items()])
    _sheet(workbook, "Search Strategies", ["Database", "Exact query", "Search date", "Records retrieved", "Records screened", "Records retained", "Notes"],
           [[item.database, item.exact_query, item.search_date, item.records_retrieved,
             item.records_screened, item.records_retained, item.notes] for item in project.searches])
    _sheet(workbook, "Imported Records", ["Record ID", "Title", "Authors", "Year", "Journal", "Abstract", "DOI", "PMID", "URL", "Keywords", "Source", "Publication type", "Active", "Import batch", "Raw imported metadata (JSON)"],
           [[r.id, r.title, r.authors, r.year, r.journal, r.abstract, r.doi, r.pmid, r.url, r.keywords, r.source_database, r.publication_type, r.active, r.import_batch_id, json.dumps(r.raw_metadata, ensure_ascii=False, sort_keys=True)] for r in project.records])
    _sheet(workbook, "Deduplication Log", ["Candidate ID", "Record A", "Record B", "Classification", "Similarity", "Reasons", "Resolution", "Reviewer", "Resolved at"],
           [[d.id, d.record_a_id, d.record_b_id, d.classification, d.similarity_score, d.reasons, d.resolution, d.reviewer, d.resolved_at] for d in project.duplicate_candidates])
    title_decisions = [d for d in project.screening_decisions if d.stage == "title_abstract"]
    full_decisions = [d for d in project.screening_decisions if d.stage == "full_text"]
    decision_headers = ["Decision ID", "Record ID", "Decision", "Reason", "Notes", "Reviewer", "Timestamp"]
    _sheet(workbook, "Screening Decisions", decision_headers,
           [[d.id, d.record_id, d.decision, d.reason, d.notes, d.reviewer, d.decided_at] for d in title_decisions])
    _sheet(workbook, "Full Text Decisions", decision_headers,
           [[d.id, d.record_id, d.decision, d.reason, d.notes, d.reviewer, d.decided_at] for d in full_decisions])
    full_current = current_decisions(project, "full_text")
    included = [r for r in project.records if r.active and full_current.get(r.id) and full_current[r.id].decision == "include"]
    _sheet(workbook, "Included Studies", ["Record ID", "Title", "Authors", "Year", "Journal", "DOI"],
           [[r.id, r.title, r.authors, r.year, r.journal, r.doi] for r in included])
    extraction_keys = list(dict.fromkeys(key for entry in project.extraction_entries for key in entry.values))
    _sheet(workbook, "Data Extraction", ["Record ID", "Template"] + extraction_keys,
           [[e.record_id, e.template] + [e.values.get(key) for key in extraction_keys] for e in project.extraction_entries])
    quality_rows = [[q.record_id, q.framework, domain.domain, domain.rating, domain.notes, q.reviewer, q.updated_at]
                    for q in project.quality_assessments for domain in q.domains]
    _sheet(workbook, "Quality Assessment", ["Record ID", "Framework", "Domain", "Rating", "Notes", "Reviewer", "Updated at"], quality_rows)
    evidence = build_evidence_table(project)
    _sheet(workbook, "Evidence Table", [column.replace("_", " ").title() for column in evidence.columns],
           [[row.get(column) for column in evidence.columns] for row in evidence.rows])
    prisma = derive_prisma_counts(project)
    prisma_fields = ["identified", "database_records", "other_source_records", "duplicates_removed", "records_screened", "records_excluded", "reports_sought", "reports_not_retrieved", "full_texts_assessed", "full_texts_excluded", "studies_included", "pending_title_abstract", "pending_full_text", "reconciled"]
    _sheet(workbook, "PRISMA Counts", ["Measure", "Count / status"], [[name.replace("_", " ").title(), getattr(prisma, name)] for name in prisma_fields])
    _sheet(workbook, "Synthesis Summary", ["Field", "Value"], [["Researcher synthesis notes", project.synthesis_notes]])
    _sheet(workbook, "Metadata", ["Field", "Value"], [
        ["Project name", project.metadata.name], ["Schema version", project.schema_version],
        ["Created at", project.metadata.created_at], ["Updated at", project.metadata.updated_at],
        ["Demonstration data", project.metadata.demonstration],
        ["Scientific integrity", "Raw imported metadata is preserved in Imported Records; decisions remain human-controlled; missing values are not treated as zero."],
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
